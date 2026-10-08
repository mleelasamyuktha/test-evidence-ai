import json
import os
from io import BytesIO
from typing import Any

import httpx
from docx import Document

from llm.llm_client import LLMClient


def read_document(data: bytes, filename: str) -> str:
    if filename.lower().endswith(".docx"):
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_rows = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + table_rows)
    return data.decode("utf-8-sig")


def build_api_test_prompt(document_text: str, openapi_contract: dict[str, Any]) -> str:
    return f"""
You are an API test automation assistant.

Read the API requirements document below and the live OpenAPI contract. Generate executable API test cases.

DOCUMENT:
{document_text}

OPENAPI CONTRACT:
{json.dumps(openapi_contract, indent=2)}

Rules:
- Use only paths and methods present in the OpenAPI contract.
- Do not use bearer tokens or authentication headers.
- Include positive, validation, not-found, duplicate, update, and delete coverage when supported by the requirements.
- For write tests, use a unique customer_id and include a Cleanup flag set to true so the runner deletes created records.
- Order dependent tests as create, get/update, then delete for the same customer_id.
- Do not add response-body assertions to a 204 No Content test.
- Use Expected or Value when an assertion compares a concrete value.
- Return only valid JSON, with no markdown fences.

Return exactly:
{{
  "SuiteName": "string",
  "TestCases": [
    {{
      "TestCaseID": "API-001",
      "Description": "string",
      "Method": "GET|POST|PUT|DELETE",
      "Path": "/api/v1/customers",
      "Body": {{}},
      "ExpectedStatus": 200,
      "Assertions": [
        {{"Path": "status", "Operator": "exists"}},
        {{"Path": "body.total_customers", "Operator": "exists"}}
      ],
      "Cleanup": false
    }}
  ]
}}
"""


def generate_api_tests(
    document_text: str,
    openapi_contract: dict[str, Any],
    model: str,
    use_context: bool = True,
) -> dict[str, Any]:
    client = LLMClient(model=model)
    response = client.generate(
        build_api_test_prompt(document_text, openapi_contract),
        use_context=use_context,
    )
    return json.loads(response)


def _read_path(payload: Any, path: str) -> Any:
    if not path or path == "body":
        return payload

    current = payload
    parts = path.split(".")
    for position, part in enumerate(parts):
        if part.startswith("[") and part.endswith("]") and part[1:-1].isdigit():
            current = current[int(part[1:-1])]
            continue
        if "[" in part and part.endswith("]"):
            field, index = part.split("[", 1)
            index = index[:-1]
            if field:
                current = current.get(field) if isinstance(current, dict) else None
            if index == "*":
                remaining = ".".join(parts[position + 1:])
                if not isinstance(current, list):
                    return None
                return [_read_path(item, remaining) if remaining else item for item in current]
            current = current[int(index)]
            continue
        if isinstance(current, dict):
            current = current.get(part)
        elif isinstance(current, list) and part.isdigit():
            current = current[int(part)]
        else:
            return None
    return current


def _assert_response(response: httpx.Response, assertion: dict[str, Any]) -> tuple[bool, str]:
    try:
        body = response.json()
    except ValueError:
        body = response.text

    path = assertion.get("Path", "status")
    if path == "status":
        actual = response.status_code
    elif path == "body":
        actual = body
    elif path.startswith("body."):
        actual = _read_path(body, path[5:])
    elif path.startswith("body["):
        actual = _read_path(body, path[4:])
    else:
        actual = _read_path(body, path)
    operator = str(assertion.get("Operator", "exists")).strip().lower()
    operator = operator.replace("-", "_").replace(" ", "_")
    operator_aliases = {
        "value": "equals",
        "equal": "equals",
        "eq": "equals",
        "all_equal": "all_equals",
        "allequal": "all_equals",
        "allequals": "all_equals",
        "allvalues_equal": "all_equals",
        "isarray": "is_array",
        "array": "is_array",
        "notexists": "not_exists",
        "minlength": "min_length",
        "maxlength": "max_length",
    }
    operator = operator_aliases.get(operator, operator)
    expected = assertion.get("Expected", assertion.get("Value"))

    if operator == "exists":
        passed = actual is not None
    elif operator == "not_exists":
        passed = actual is None
    elif operator == "is_array":
        passed = isinstance(actual, list)
    elif operator in {"all_equal", "all_equals"}:
        passed = isinstance(actual, list) and len(set(str(value) for value in actual)) <= 1
    elif operator == "equals":
        if expected is None:
            passed = actual is not None
        elif isinstance(actual, list):
            passed = all(value == expected for value in actual)
        else:
            passed = actual == expected
    elif operator == "contains":
        passed = expected in actual
    elif operator == "min_length":
        passed = actual is not None and expected is not None and len(actual) >= int(expected)
    elif operator == "max_length":
        passed = actual is not None and expected is not None and len(actual) <= int(expected)
    elif operator == "length":
        passed = actual is not None and expected is not None and len(actual) == int(expected)
    else:
        return False, f"Unsupported assertion operator: {operator}"
    return passed, f"{path}: actual={actual!r}, operator={operator}, expected={expected!r}"


def run_api_tests(suite: dict[str, Any], base_url: str) -> list[dict[str, Any]]:
    results = []
    created_ids = []

    with httpx.Client(base_url=base_url, timeout=30.0) as client:
        for test_case in suite.get("TestCases", []):
            method = test_case["Method"].upper()
            path = test_case["Path"]
            body = test_case.get("Body") or None
            try:
                response = client.request(method, path, json=body)
                if method == "POST" and response.status_code == 201 and test_case.get("Cleanup"):
                    created_id = response.json().get("customer_id")
                    if created_id is not None:
                        created_ids.append(created_id)

                checks = [_assert_response(response, assertion) for assertion in test_case.get("Assertions", [])]
                passed = response.status_code == test_case.get("ExpectedStatus") and all(item[0] for item in checks)
                messages = list(dict.fromkeys(item[1] for item in checks))
                message = "; ".join(messages) or f"HTTP {response.status_code}"
                results.append({
                    "TestCaseID": test_case["TestCaseID"],
                    "Description": test_case.get("Description", ""),
                    "Status": "PASS" if passed else "FAIL",
                    "Message": message,
                    "Evidence": {"Method": method, "Path": path, "ExpectedStatus": test_case.get("ExpectedStatus"), "ActualStatus": response.status_code, "Response": response.json() if response.content and response.text.strip() else None},
                })
            except Exception as error:
                results.append({
                    "TestCaseID": test_case["TestCaseID"],
                    "Description": test_case.get("Description", ""),
                    "Status": "ERROR",
                    "Message": str(error),
                    "Evidence": {"Method": method, "Path": path},
                })

        for customer_id in created_ids:
            client.delete(f"/api/v1/customers/{customer_id}")

    return results